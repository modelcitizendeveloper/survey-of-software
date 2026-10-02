"""L4 — the spreadsheet engines for survey 1.057.

Settles S1's first open question: is calamine faster than openpyxl, and by how much on
what. Two axes:

  DIRECT   openpyxl (normal and read_only) vs python-calamine, reading the same workbook.
  VIA PANDAS  read_excel through each engine pandas registers that can open .xlsx —
              openpyxl and calamine — so the cost of the delegation is visible too.

Also records what each engine returns for the spreadsheet-only hazards in hostile.xlsx: a
merged header above the real header row, a text-formatted column holding digits, and a
date cell whose serial differs from its display.

Output: results/l4-spreadsheet.json
"""

from __future__ import annotations

import json
import pathlib
import subprocess
import sys
import textwrap

PY = str(pathlib.Path(".venv/bin/python").absolute())
ROUNDS = 3

BODIES = {
    "openpyxl": """
import openpyxl
wb = openpyxl.load_workbook(PATH)
ws = wb[SHEET]
rows = list(ws.iter_rows(values_only=True))
n = len(rows)
""",
    "openpyxl (read_only)": """
import openpyxl
wb = openpyxl.load_workbook(PATH, read_only=True)
ws = wb[SHEET]
rows = list(ws.iter_rows(values_only=True))
n = len(rows)
wb.close()
""",
    "python-calamine": """
from python_calamine import CalamineWorkbook
wb = CalamineWorkbook.from_path(PATH)
rows = wb.get_sheet_by_name(SHEET).to_python()
n = len(rows)
""",
    "pandas via openpyxl": """
import pandas as pd
df = pd.read_excel(PATH, sheet_name=SHEET, engine='openpyxl', header=None)
n = len(df)
""",
    "pandas via calamine": """
import pandas as pd
df = pd.read_excel(PATH, sheet_name=SHEET, engine='calamine', header=None)
n = len(df)
""",
}

RUNNER = """
import json, resource, time
PATH = {path!r}; SHEET = {sheet!r}
best = None
for _ in range({rounds}):
    _t0 = time.perf_counter()
{body}
    dt = time.perf_counter() - _t0
    best = dt if best is None else min(best, dt)
peak_kb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
print(json.dumps({{"seconds": round(best, 4), "rows": n,
                   "peak_rss_mb": round(peak_kb / 1024, 1)}}))
"""

# What each engine hands back for the three spreadsheet-only hazards.
HAZARD_PROBE = """
import json
PATH = {path!r}
out = {{}}
try:
{body}
except Exception as e:
    out['error'] = f'{{type(e).__name__}}: {{e}}'
print(json.dumps(out, default=str))
"""

PROBES = {
    "openpyxl": """
    import openpyxl
    wb = openpyxl.load_workbook(PATH)
    ws = wb['data']
    out['a1_merged_header'] = ws.cell(1, 1).value
    out['header_row_2'] = [ws.cell(2, c).value for c in range(1, 4)]
    out['zero_padded_row3'] = ws.cell(5, 2).value        # data row 3 -> sheet row 5
    out['date_cell'] = ws.cell(3, 6).value               # data row 1, 'when'
    out['sheet_names'] = wb.sheetnames
""",
    "python-calamine": """
    from python_calamine import CalamineWorkbook
    wb = CalamineWorkbook.from_path(PATH)
    rows = wb.get_sheet_by_name('data').to_python()
    out['a1_merged_header'] = rows[0][0]
    out['header_row_2'] = rows[1][:3]
    out['zero_padded_row3'] = rows[4][1]
    out['date_cell'] = rows[2][5]
    out['sheet_names'] = wb.sheet_names
""",
    "pandas via openpyxl": """
    import pandas as pd
    df = pd.read_excel(PATH, sheet_name='data', engine='openpyxl', header=None)
    out['a1_merged_header'] = df.iloc[0, 0]
    out['header_row_2'] = list(df.iloc[1, :3])
    out['zero_padded_row3'] = df.iloc[4, 1]
    out['date_cell'] = df.iloc[2, 5]
""",
    "pandas via calamine": """
    import pandas as pd
    df = pd.read_excel(PATH, sheet_name='data', engine='calamine', header=None)
    out['a1_merged_header'] = df.iloc[0, 0]
    out['header_row_2'] = list(df.iloc[1, :3])
    out['zero_padded_row3'] = df.iloc[4, 1]
    out['date_cell'] = df.iloc[2, 5]
""",
    "sheet 0 BY INDEX (the hazard)": """
    import pandas as pd
    from python_calamine import CalamineWorkbook
    # Takes the first sheet by position, which is what read_excel does by default and what
    # a reader who has not opened the file assumes is the table.
    df = pd.read_excel(PATH, sheet_name=0, header=None)
    out['first_cell'] = df.iloc[0, 0]
    out['shape'] = list(df.shape)
    wb = CalamineWorkbook.from_path(PATH)
    out['sheet_0_name'] = wb.sheet_names[0]
""",
    "pandas default engine": """
    import pandas as pd
    df = pd.read_excel(PATH, sheet_name='data', header=None)
    out['a1_merged_header'] = df.iloc[0, 0]
    out['zero_padded_row3'] = df.iloc[4, 1]
    out['engine_used'] = 'not reported by pandas; this is whatever it chose'
""",
}


def make_clean_xlsx(path: pathlib.Path, rows: int) -> None:
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(path), {"constant_memory": True})
    ws = wb.add_worksheet("data")
    cols = ["row_id", "name", "amount", "note_text", "when"]
    for c, name in enumerate(cols):
        ws.write_string(0, c, name)
    for i in range(1, rows + 1):
        ws.write_string(i, 0, str(i))
        ws.write_string(i, 1, f"Person {i}")
        ws.write_string(i, 2, f"{(i*37)%10000}.{i%100:02d}")
        ws.write_string(i, 3, f"note {i}")
        ws.write_string(i, 4, "2026-01-15")
    wb.close()


def run(label: str, path: str, sheet: str, rounds: int) -> dict:
    body = "\n".join("    " + ln for ln in BODIES[label].strip().splitlines())
    code = RUNNER.format(path=path, sheet=sheet, rounds=rounds, body=body)
    out = subprocess.run([PY, "-c", code], capture_output=True, text=True)
    if out.returncode != 0:
        # A child killed by a signal (OOM, returncode -9) writes NOTHING to stderr, and
        # [-1] on an empty list raises IndexError from inside the error handler — losing
        # the whole run instead of recording one failure.
        err = out.stderr.strip().splitlines()
        return {"error": (err[-1][:200] if err else
                          f"no stderr; returncode {out.returncode}")}
    return json.loads(out.stdout.strip())


def probe(label: str, path: str) -> dict:
    # The PROBES bodies are written already indented for readability, so dedent BEFORE
    # re-indenting — otherwise line 1 gets 4 spaces and the rest get 8.
    body = "\n".join("    " + ln for ln in
                     textwrap.dedent(PROBES[label]).strip().splitlines())
    code = HAZARD_PROBE.format(path=path, body=body)
    out = subprocess.run([PY, "-c", code], capture_output=True, text=True)
    if out.returncode != 0:
        # A child killed by a signal (OOM, returncode -9) writes NOTHING to stderr, and
        # [-1] on an empty list raises IndexError from inside the error handler — losing
        # the whole run instead of recording one failure.
        err = out.stderr.strip().splitlines()
        return {"error": (err[-1][:200] if err else
                          f"no stderr; returncode {out.returncode}")}
    return json.loads(out.stdout.strip())


def main() -> int:
    fx = pathlib.Path("fixtures")
    sizes = {}
    for n in (5_000, 50_000):
        p = fx / f"clean_{n}.xlsx"
        if not p.exists():
            make_clean_xlsx(p, n)
        sizes[f"{n//1000}k"] = str(p)

    results = {"rounds": ROUNDS, "timing": {}, "hazards": {}}

    for size, path in sizes.items():
        mb = pathlib.Path(path).stat().st_size / 1e6
        print(f"\n=== {size} rows ({mb:.1f} MB xlsx) ===")
        per = {}
        for label in BODIES:
            r = run(label, path, "data", ROUNDS)
            per[label] = r
            if "error" in r:
                print(f"  {label:24} ERROR {r['error'][:70]}")
            else:
                print(f"  {label:24} {r['seconds']:7.3f} s   peak {r['peak_rss_mb']:7.1f} MB   "
                      f"{r['rows']:,} rows")
        results["timing"][size] = per

    print("\n=== spreadsheet-only hazards (hostile.xlsx) ===")
    hp = str(fx / "hostile.xlsx")
    for label in PROBES:
        r = probe(label, hp)
        results["hazards"][label] = r
        print(f"  {label}")
        for k, v in r.items():
            print(f"      {k:22} {v!r}")

    out = pathlib.Path("results"); out.mkdir(exist_ok=True)
    (out / "l4-spreadsheet.json").write_text(json.dumps(results, indent=2, default=str))
    print("\nwrote results/l4-spreadsheet.json")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
